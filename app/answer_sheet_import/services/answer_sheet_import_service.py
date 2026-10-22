# -*- coding: utf-8 -*-
"""
Pré-visualização da importação de cartão-resposta (gabarito).

Lê planilha (CSV/XLSX/XLS) ou texto colado do Excel, resolve disciplinas e habilidades
e devolve um rascunho validado. Nada é persistido: o frontend preenche o cadastro e o
gabarito continua sendo criado por POST /answer-sheets/create-gabaritos.
"""

from __future__ import annotations

import csv
import io
import logging
import re
import uuid
from typing import Any, Callable, Dict, Iterable, List, Optional, Sequence, Set, Tuple

from openpyxl import load_workbook
from sqlalchemy import func
from xlrd import open_workbook

from app.models.grades import Grade
from app.models.skill import Skill
from app.models.subject import Subject
from app.question_import.importer import _is_uuid, _norm_name
from app.services.skills_map_service import normalize_skill_code

logger = logging.getLogger(__name__)

ALLOWED_EXTENSIONS = {"csv", "xlsx", "xls"}
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 200_000
MAX_SHEET_ROWS = 500
MAX_QUESTIONS = 104
MAX_SUBJECT_BLOCKS = 4
MIN_ALTERNATIVES = 2
ALTERNATIVE_LETTERS = "ABCDE"
METADATA_MAX_LEN = 200

REQUIRED_COLUMNS = ("question", "answer", "subject")
POSITIONAL_COLUMNS = ("question", "answer", "subject", "skill", "alternatives")
COLUMN_LABELS = {"question": "Questão", "answer": "Resposta", "subject": "Disciplina"}

_COLUMN_ALIASES = {
    "question": ("questao", "questoes", "q", "n", "numero", "num"),
    "answer": ("resposta", "gabarito", "correta", "resposta correta", "alternativa correta"),
    "subject": ("disciplina", "componente", "componente curricular", "materia"),
    "skill": ("habilidade", "habilidades", "codigo da habilidade", "codigo habilidade", "descritor", "bncc"),
    "alternatives": (
        "alternativas", "opcoes", "qtd alternativas", "qtd de alternativas", "quantidade de alternativas",
        "n alternativas", "n de alternativas", "numero de alternativas",
    ),
}
_COLUMN_BY_ALIAS = {alias: key for key, aliases in _COLUMN_ALIASES.items() for alias in aliases}

_METADATA_ALIASES = {
    "title": ("titulo", "titulo da avaliacao", "avaliacao"),
    "institution": ("secretaria", "instituicao"),
    "grades": ("serie", "series", "ano", "anos", "serie ano", "series anos"),
}
_METADATA_BY_ALIAS = {alias: key for key, aliases in _METADATA_ALIASES.items() for alias in aliases}
_METADATA_LABELS = {"title": "Título", "institution": "Secretaria"}

_QUESTION_RE = re.compile(r"(?:q|questao)?\s*(\d{1,3})\s*[).]?")
_ANSWER_RE = re.compile(r"\(?\s*([a-z])\s*\)?\s*[).]?")
_SKILL_TOKEN_RE = re.compile(r"[A-Z0-9][A-Z0-9._]*")
_SKILL_HEAD_SPLIT_RE = re.compile(r"\s+[-–]\s+|\s*:\s*")


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _norm_header(value: Any) -> str:
    text = _norm_name(_cell_text(value).replace("\ufeff", ""))
    text = re.sub(r"[^a-z0-9 ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _column_key(header: str) -> Optional[str]:
    if not header:
        return None
    key = _COLUMN_BY_ALIAS.get(header)
    if key:
        return key
    first_word = header.split(" ")[0]
    return _COLUMN_BY_ALIAS.get(first_word) if len(first_word) >= 4 else None


def _non_empty_rows(rows: Iterable[Sequence[Any]]) -> List[Tuple[int, List[Any]]]:
    """Mantém o número da linha original (1-based) para mensagens de erro."""
    result: List[Tuple[int, List[Any]]] = []
    for line, row in enumerate(rows, start=1):
        cells = list(row or [])
        if any(_cell_text(cell) for cell in cells):
            result.append((line, cells))
            if len(result) > MAX_SHEET_ROWS:
                raise ValueError(f"A planilha excede o limite de {MAX_SHEET_ROWS} linhas preenchidas.")
    return result


def _decode_text(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("latin-1")


def _parse_delimited_text(content: str) -> List[List[str]]:
    sample = content[:4096]
    if "\t" in sample:
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=";,").delimiter
        except csv.Error:
            first_line = sample.splitlines()[0] if sample else ""
            delimiter = ";" if ";" in first_line else ","
    return list(csv.reader(io.StringIO(content), delimiter=delimiter))


def _split_sheets(sheets: Dict[str, List[Sequence[Any]]]) -> Tuple[List[Sequence[Any]], List[Sequence[Any]]]:
    """Aba 'Gabarito' (ou a primeira que não seja 'Dados') + aba opcional 'Dados'."""
    named = {_norm_header(name): rows for name, rows in sheets.items()}
    answer_rows = named.get("gabarito")
    if answer_rows is None:
        answer_rows = next((rows for name, rows in sheets.items() if _norm_header(name) != "dados"), [])
    return answer_rows, named.get("dados", [])


def _read_file_tables(file_storage) -> Tuple[List[Sequence[Any]], List[Sequence[Any]]]:
    filename = (getattr(file_storage, "filename", "") or "").strip()
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError("Formato de arquivo não suportado. Use CSV, XLSX ou XLS.")

    content = file_storage.read(MAX_FILE_BYTES + 1)
    if not content:
        raise ValueError("Arquivo vazio.")
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("Arquivo excede o limite de 5 MB.")

    if extension == "csv":
        return _parse_delimited_text(_decode_text(content)), []

    try:
        if extension == "xlsx":
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                sheets = {
                    ws.title: list(ws.iter_rows(values_only=True))
                    for ws in workbook.worksheets
                }
            finally:
                workbook.close()
        else:
            book = open_workbook(file_contents=content)
            sheets = {sh.name: [sh.row_values(i) for i in range(sh.nrows)] for sh in book.sheets()}
    except Exception as exc:
        logger.warning("Falha ao ler planilha de cartão-resposta: %s", exc)
        raise ValueError("Não foi possível ler o arquivo. Verifique se é uma planilha válida.") from exc

    return _split_sheets(sheets)


def _map_columns(first_cells: Sequence[Any]) -> Tuple[Dict[str, int], bool]:
    """Retorna (coluna -> índice, possui_cabeçalho). Sem cabeçalho, usa a ordem posicional."""
    mapping: Dict[str, int] = {}
    for idx, cell in enumerate(first_cells):
        key = _column_key(_norm_header(cell))
        if key and key not in mapping:
            mapping[key] = idx

    if mapping:
        missing = [COLUMN_LABELS[col] for col in REQUIRED_COLUMNS if col not in mapping]
        if missing:
            raise ValueError(
                f"Colunas obrigatórias ausentes: {', '.join(missing)}. "
                "Use os cabeçalhos Questão, Resposta e Disciplina."
            )
        return mapping, True

    if _parse_question_number(_cell_text(first_cells[0]) if first_cells else "") is None:
        raise ValueError(
            "Cabeçalho não reconhecido. A primeira linha deve conter Questão, Resposta e Disciplina."
        )
    return {key: idx for idx, key in enumerate(POSITIONAL_COLUMNS)}, False


def _parse_question_number(text: str) -> Optional[int]:
    match = _QUESTION_RE.fullmatch(_norm_name(text))
    if not match:
        return None
    number = int(match.group(1))
    return number if number >= 1 else None


def _parse_answer(text: str) -> Optional[str]:
    match = _ANSWER_RE.fullmatch(_norm_name(text))
    if not match:
        return None
    letter = match.group(1).upper()
    return letter if letter in ALTERNATIVE_LETTERS else None


def _parse_alternatives(text: str) -> Tuple[Optional[int], Optional[str]]:
    """Aceita número (ex.: 4) ou letras consecutivas a partir de A (ex.: A,B,C)."""
    if not text:
        return None, None
    max_count = len(ALTERNATIVE_LETTERS)
    range_error = f"Quantidade de alternativas deve estar entre {MIN_ALTERNATIVES} e {max_count}."

    if re.fullmatch(r"\d+", text):
        count = int(text)
    else:
        letters = "".join(re.findall(r"[A-Za-z]", text)).upper()
        if len(letters) > max_count:
            return None, range_error
        if not letters or letters != ALTERNATIVE_LETTERS[: len(letters)]:
            return None, "Alternativas devem ser consecutivas a partir de A (ex.: A,B,C) ou um número."
        count = len(letters)

    if count < MIN_ALTERNATIVES or count > max_count:
        return None, range_error
    return count, None


def _build_subject_resolver() -> Callable[[str], Optional[Subject]]:
    """Mesmo critério do importador de questões: id → nome exato → nome contido."""
    subjects = Subject.query.all()
    by_id = {str(s.id): s for s in subjects}
    by_name = [(_norm_name(s.name or ""), s) for s in subjects]
    cache: Dict[str, Optional[Subject]] = {}

    def resolve(raw: str) -> Optional[Subject]:
        key = raw.strip()
        if key not in cache:
            subject = by_id.get(key) if _is_uuid(key) else None
            if subject is None:
                needle = _norm_name(key)
                subject = next((s for name, s in by_name if name == needle), None) or next(
                    (s for name, s in by_name if needle and needle in name), None
                )
            cache[key] = subject
        return cache[key]

    return resolve


def _skill_code_candidates(raw: str) -> List[str]:
    """
    Detecta o código da habilidade no texto da célula:
    'EF05MA01', 'ef05ma01', 'EF05MA01 - Descrição', 'Descrição (EF05MA01)'.
    """
    text = normalize_skill_code(raw)
    head = _SKILL_HEAD_SPLIT_RE.split(text, maxsplit=1)[0].strip()
    candidates = [text, head, head.replace(" ", "")]
    candidates.extend(
        token
        for token in _SKILL_TOKEN_RE.findall(text)
        if len(token) >= 3 and re.search(r"[A-Z]", token) and re.search(r"\d", token)
    )
    return list(dict.fromkeys(c for c in candidates if c))


def _load_skills(raw_values: Iterable[str]) -> Tuple[Dict[str, List[Skill]], Dict[str, Skill]]:
    codes: Set[str] = set()
    ids: Set[uuid.UUID] = set()
    for raw in raw_values:
        if _is_uuid(raw):
            ids.add(uuid.UUID(raw.strip()))
        else:
            codes.update(_skill_code_candidates(raw))

    by_code: Dict[str, List[Skill]] = {}
    if codes:
        rows = Skill.query.filter(func.upper(func.trim(Skill.code)).in_(sorted(codes))).all()
        for skill in rows:
            by_code.setdefault(normalize_skill_code(skill.code), []).append(skill)

    by_id: Dict[str, Skill] = {}
    if ids:
        by_id = {str(s.id): s for s in Skill.query.filter(Skill.id.in_(list(ids))).all()}
    return by_code, by_id


def _pick_skill(options: List[Skill], subject_id: Optional[str], grade_ids: Set[str]) -> Skill:
    """Mesmo código pode existir em mais de uma disciplina/série: prioriza disciplina e série do cartão."""
    def rank(skill: Skill) -> Tuple[bool, bool]:
        same_subject = bool(subject_id) and str(skill.subject_id or "") == str(subject_id)
        same_grade = any(str(g.id) in grade_ids for g in (skill.grades or []))
        return same_subject, same_grade

    return max(options, key=rank)


def _resolve_skill(
    raw: str,
    subject: Optional[Subject],
    by_code: Dict[str, List[Skill]],
    by_id: Dict[str, Skill],
    grade_ids: Set[str],
) -> Optional[Skill]:
    if _is_uuid(raw):
        return by_id.get(str(uuid.UUID(raw.strip())))
    for code in _skill_code_candidates(raw):
        options = by_code.get(code)
        if options:
            return _pick_skill(options, str(subject.id) if subject else None, grade_ids)
    return None


def _norm_grade_name(value: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[ºª°]", "", _norm_name(value))).strip()


def _clean_grade_ids(values: Any) -> List[str]:
    if isinstance(values, str):
        values = [values]
    if not isinstance(values, (list, tuple)):
        return []
    return [str(v).strip() for v in values if _is_uuid(v)]


def _parse_metadata(rows: List[Tuple[int, List[Any]]]) -> Tuple[Dict[str, Any], List[str]]:
    """Aba 'Dados' no formato chave/valor: Título, Secretaria, Séries (separadas por ; ou ,)."""
    metadata: Dict[str, Any] = {"title": None, "institution": None, "grades": []}
    warnings: List[str] = []
    grade_names: List[str] = []

    for _, cells in rows:
        key = _METADATA_BY_ALIAS.get(_norm_header(cells[0] if cells else None))
        value = _cell_text(cells[1]) if len(cells) > 1 else ""
        if not key or not value:
            continue
        if key == "grades":
            grade_names.extend(part.strip() for part in re.split(r"[;,]", value) if part.strip())
        else:
            if len(value) > METADATA_MAX_LEN:
                value = value[:METADATA_MAX_LEN]
                warnings.append(f"{_METADATA_LABELS[key]} cortado para {METADATA_MAX_LEN} caracteres.")
            metadata[key] = value

    if grade_names:
        by_name = {_norm_grade_name(g.name or ""): g for g in Grade.query.all()}
        for name in grade_names:
            grade = by_name.get(_norm_grade_name(name))
            if grade is None:
                warnings.append(f"Série não encontrada: {name}.")
            elif all(item["id"] != str(grade.id) for item in metadata["grades"]):
                metadata["grades"].append({"id": str(grade.id), "name": grade.name})

    return metadata, warnings


def _build_blocks(rows: List[Dict[str, Any]]) -> Tuple[List[Dict[str, Any]], List[str]]:
    """Agrupa questões consecutivas da mesma disciplina em blocos."""
    blocks: List[Dict[str, Any]] = []
    for row in rows:
        subject, question = row["subject"], row["question"]
        if subject is None or question is None:
            return [], []
        if blocks and blocks[-1]["subject_id"] == subject["id"]:
            blocks[-1]["end_question"] = question
            blocks[-1]["questions_count"] += 1
        else:
            blocks.append({
                "subject_id": subject["id"],
                "subject_name": subject["name"],
                "start_question": question,
                "end_question": question,
                "questions_count": 1,
            })

    errors: List[str] = []
    subject_ids = [b["subject_id"] for b in blocks]
    repeated = sorted({b["subject_name"] for b in blocks if subject_ids.count(b["subject_id"]) > 1})
    if repeated:
        errors.append(
            "As questões de cada disciplina devem estar em sequência: " + ", ".join(repeated) + "."
        )
    if len(blocks) > MAX_SUBJECT_BLOCKS:
        errors.append(f"Máximo de {MAX_SUBJECT_BLOCKS} disciplinas (blocos) por cartão-resposta.")
    return blocks, errors


def _build_preview(
    answer_rows: List[Sequence[Any]],
    metadata_rows: List[Sequence[Any]],
    grade_ids: Any,
) -> Dict[str, Any]:
    rows = _non_empty_rows(answer_rows)
    if not rows:
        raise ValueError("Nenhuma questão encontrada.")
    mapping, has_header = _map_columns(rows[0][1])
    data_rows = rows[1:] if has_header else rows
    if not data_rows:
        raise ValueError("Nenhuma questão encontrada abaixo do cabeçalho.")
    if len(data_rows) > MAX_QUESTIONS:
        raise ValueError(f"Máximo de {MAX_QUESTIONS} questões por cartão-resposta.")

    metadata, warnings = _parse_metadata(_non_empty_rows(metadata_rows))
    grade_id_set = set(_clean_grade_ids(grade_ids)) | {g["id"] for g in metadata["grades"]}

    def cell(cells: List[Any], key: str) -> str:
        idx = mapping.get(key)
        return _cell_text(cells[idx]) if idx is not None and idx < len(cells) else ""

    raw_rows = [
        {"line": line, **{key: cell(cells, key) for key in POSITIONAL_COLUMNS}}
        for line, cells in data_rows
    ]
    resolve_subject = _build_subject_resolver()
    skills_by_code, skills_by_id = _load_skills(r["skill"] for r in raw_rows if r["skill"])

    parsed: List[Dict[str, Any]] = []
    seen: Dict[int, int] = {}
    for raw in raw_rows:
        errors: List[str] = []
        row_warnings: List[str] = []

        question = _parse_question_number(raw["question"])
        if question is None:
            errors.append("Número da questão inválido.")
        elif question in seen:
            errors.append(f"Questão {question} repetida (linha {seen[question]}).")
        else:
            seen[question] = raw["line"]

        answer = _parse_answer(raw["answer"])
        if answer is None:
            errors.append(
                f"Resposta deve ser uma letra de A a {ALTERNATIVE_LETTERS[-1]}." if raw["answer"]
                else "Resposta obrigatória."
            )

        alternatives_count, alternatives_error = _parse_alternatives(raw["alternatives"])
        if alternatives_error:
            errors.append(alternatives_error)

        subject = resolve_subject(raw["subject"]) if raw["subject"] else None
        if subject is None:
            errors.append(
                f"Disciplina não encontrada: {raw['subject']}." if raw["subject"] else "Disciplina obrigatória."
            )

        skill = None
        if raw["skill"]:
            skill = _resolve_skill(raw["skill"], subject, skills_by_code, skills_by_id, grade_id_set)
            if skill is None:
                row_warnings.append(f"Habilidade não encontrada: {raw['skill']}.")
            elif subject and skill.subject_id and str(skill.subject_id) != str(subject.id):
                row_warnings.append(f"Habilidade {skill.code} está cadastrada em outra disciplina.")

        parsed.append({
            "line": raw["line"],
            "question": question,
            "answer": answer,
            "alternatives_count": alternatives_count,
            "subject": {"id": str(subject.id), "name": subject.name} if subject else None,
            "subject_input": raw["subject"],
            "skill": (
                {"id": str(skill.id), "code": skill.code, "description": skill.description}
                if skill else None
            ),
            "skill_input": raw["skill"],
            "errors": errors,
            "warnings": row_warnings,
        })

    parsed.sort(key=lambda r: (r["question"] is None, r["question"] or 0, r["line"]))

    general_errors: List[str] = []
    if seen:
        missing = [n for n in range(1, max(seen) + 1) if n not in seen]
        if missing:
            shown = ", ".join(str(n) for n in missing[:20])
            general_errors.append(
                f"A numeração deve ser sequencial a partir de 1. Faltando: {shown}"
                + ("..." if len(missing) > 20 else ".")
            )

    blocks, block_errors = _build_blocks(parsed)
    general_errors.extend(block_errors)

    return {
        "rows": parsed,
        "blocks": blocks,
        "metadata": metadata,
        "errors": general_errors,
        "warnings": warnings,
        "summary": {
            "total_questions": len(parsed),
            "rows_with_errors": sum(1 for r in parsed if r["errors"]),
            "skills_detected": sum(1 for r in parsed if r["skill"]),
        },
    }


def preview_from_file(file_storage, grade_ids: Any = None) -> Dict[str, Any]:
    answer_rows, metadata_rows = _read_file_tables(file_storage)
    return _build_preview(answer_rows, metadata_rows, grade_ids)


def preview_from_text(text: Any, grade_ids: Any = None) -> Dict[str, Any]:
    if not isinstance(text, str) or not text.strip():
        raise ValueError("Cole o conteúdo copiado do Excel.")
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError("O conteúdo colado excede o limite permitido.")
    return _build_preview(_parse_delimited_text(text), [], grade_ids)
