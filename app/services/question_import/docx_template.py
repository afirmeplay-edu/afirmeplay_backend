# -*- coding: utf-8 -*-
"""Gera o template DOCX para importação em lote de questões (múltipla escolha)."""

from __future__ import annotations

from io import BytesIO
from typing import Any, Dict, Optional

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt, RGBColor

from app.services.question_import.constants import ALLOWED_DIFFICULTIES


def _add_heading(doc: Document, text: str, size: int = 16) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)


def _add_muted(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)


def _add_marker(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _add_locked_meta(doc: Document, label: str, value: str) -> None:
    """Campo pré-preenchido pelo sistema — usuário não deve editar."""
    p = doc.add_paragraph()
    run = p.add_run(f"{label}: {value}")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x16, 0x65, 0x34)


def _add_user_meta(doc: Document, label: str, value: str) -> None:
    """Campo que o usuário preenche (ex.: dificuldade por questão)."""
    p = doc.add_paragraph()
    run = p.add_run(f"{label}: {value}")
    run.font.size = Pt(11)


def build_questions_import_template(context: Optional[Dict[str, Any]] = None) -> BytesIO:
    """
    Monta um .docx só para questões de múltipla escolha.

    context (obrigatório na rota):
      subjectId, subjectName, gradeId, gradeName
    Disciplina/série vêm do formulário; dificuldade é por questão no arquivo.
    """
    context = context or {}
    subject_name = context.get("subjectName") or "—"
    subject_id = context.get("subjectId") or ""
    grade_name = context.get("gradeName") or "—"
    grade_id = context.get("gradeId") or ""

    doc = Document()

    _add_heading(doc, "Template de importação de questões — Afirme Play")
    _add_muted(
        doc,
        "Este arquivo é só para questões de múltipla escolha. "
        "Disciplina e série já foram definidas no formulário (campos em verde — não altere). "
        "Não altere os marcadores em amarelo. Você pode colar imagens no enunciado, "
        "nas alternativas ou na solução.",
    )
    doc.add_paragraph()

    _add_heading(doc, "Contexto deste arquivo (não editar)", size=13)
    _add_locked_meta(doc, "Disciplina", subject_name)
    if subject_id:
        _add_locked_meta(doc, "SubjectId", str(subject_id))
    _add_locked_meta(doc, "Série", grade_name)
    if grade_id:
        _add_locked_meta(doc, "GradeId", str(grade_id))
    doc.add_paragraph()

    _add_heading(doc, "Dificuldade (obrigatória em cada questão)", size=13)
    _add_muted(
        doc,
        "Copie e cole um dos textos abaixo exatamente no campo Dificuldade de cada questão. "
        "Assim você evita erro de digitação. Variações leves de acento/maiúscula são aceitas, "
        "mas o mais seguro é copiar e colar. Cada questão pode ter um nível diferente.",
    )
    for label in ALLOWED_DIFFICULTIES:
        doc.add_paragraph(label, style="List Bullet")

    doc.add_paragraph()
    _add_heading(doc, "O que você preenche em cada questão", size=13)
    for line in [
        "Dificuldade: obrigatória (copie e cole um dos quatro textos acima)",
        "Habilidade: código BNCC/código interno (ex.: EF05MA01) — opcional",
        "Título, Comando, Número, Valor — opcionais",
        "Marque a alternativa correta com [CORRETA] no final da linha",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_paragraph()
    _add_heading(doc, "Exemplo 1", size=13)

    _add_marker(doc, "=== QUESTÃO ===")
    _add_locked_meta(doc, "Disciplina", subject_name)
    if subject_id:
        _add_locked_meta(doc, "SubjectId", str(subject_id))
    _add_locked_meta(doc, "Série", grade_name)
    if grade_id:
        _add_locked_meta(doc, "GradeId", str(grade_id))
    _add_user_meta(doc, "Dificuldade", "Adequado")
    doc.add_paragraph("Habilidade: EF05MA01")
    doc.add_paragraph("Título: Frações equivalentes")
    doc.add_paragraph()
    _add_marker(doc, "Enunciado:")
    doc.add_paragraph(
        "Qual das frações abaixo é equivalente a 1/2? "
        "(Você pode colar uma imagem aqui, se quiser.)"
    )
    doc.add_paragraph()
    _add_marker(doc, "Alternativas:")
    doc.add_paragraph("A) 2/4 [CORRETA]")
    doc.add_paragraph("B) 1/3")
    doc.add_paragraph("C) 3/5")
    doc.add_paragraph("D) 2/5")
    doc.add_paragraph()
    _add_marker(doc, "Solução:")
    doc.add_paragraph(
        "Multiplicando numerador e denominador de 1/2 por 2 obtemos 2/4, "
        "que é uma fração equivalente."
    )
    _add_marker(doc, "=== FIM ===")

    doc.add_paragraph()
    _add_heading(doc, "Exemplo 2 (outra dificuldade no mesmo arquivo)", size=13)

    _add_marker(doc, "=== QUESTÃO ===")
    _add_locked_meta(doc, "Disciplina", subject_name)
    if subject_id:
        _add_locked_meta(doc, "SubjectId", str(subject_id))
    _add_locked_meta(doc, "Série", grade_name)
    if grade_id:
        _add_locked_meta(doc, "GradeId", str(grade_id))
    _add_user_meta(doc, "Dificuldade", "Básico")
    doc.add_paragraph("Título: Números pares")
    doc.add_paragraph()
    _add_marker(doc, "Enunciado:")
    doc.add_paragraph("Qual dos números abaixo é par?")
    doc.add_paragraph()
    _add_marker(doc, "Alternativas:")
    doc.add_paragraph("A) 3")
    doc.add_paragraph("B) 7")
    doc.add_paragraph("C) 8 [CORRETA]")
    doc.add_paragraph("D) 9")
    doc.add_paragraph()
    _add_marker(doc, "Solução:")
    doc.add_paragraph("8 é divisível por 2, portanto é par.")
    _add_marker(doc, "=== FIM ===")

    doc.add_paragraph()
    _add_muted(
        doc,
        "Dica: copie o bloco inteiro (da linha === QUESTÃO === até === FIM ===) "
        "para adicionar mais questões. Mantenha Disciplina/SubjectId/Série/GradeId "
        "iguais; altere só a Dificuldade (copiar/colar) e o conteúdo da questão.",
    )

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
